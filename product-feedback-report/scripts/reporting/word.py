"""Generate an editable Word version of the tracking feedback report."""

from __future__ import annotations

import re
import urllib.request
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_ORIENT
from docx.enum.table import (
    WD_ALIGN_VERTICAL,
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_ROW_HEIGHT_RULE,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Twips

from scripts.reporting.html import (
    DeliveryReport,
    JdSale,
    ProductInfo,
    SocialReport,
    normalize_product_key,
    read_social_report,
)


ROOT = Path(__file__).resolve().parents[2]
BLACK = "000000"
HEADER_FILL = "D9D9D9"
LABEL_FILL = "E7EFFA"
TRACKED_FILL = "FCE4D6"
BORDER_COLOR = BLACK
# Match the branded Chinese font used by the HTML report and reference DOCX.
CHINESE_FONT = "方正FW筑紫黑 简 R"
LATIN_FONT = "Heytea Sans Serif"
LETTER_SPACING_TWENTIETHS_PT = "11"  # 0.02 cm ≈ 0.567 pt


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _format_count(value: float) -> str:
    return f"{value:,.0f}"


def _format_daily(value: float) -> str:
    return f"{value:.1f}"


def _format_percent(value: float) -> str:
    return f"{value:.1%}"


def _format_positive_rate(value: float) -> str:
    return f"{value:.0%}"


def _top_labels(values: tuple[tuple[str, int], ...]) -> str:
    return "、".join(label for label, _ in values) or "暂无明确高频观点"


def _safe_config_path(value: str) -> Path | None:
    if not value:
        return None
    candidate = Path(value)
    if not candidate.is_absolute():
        candidate = (ROOT / candidate).resolve()
    try:
        candidate.relative_to(ROOT.parent.resolve())
    except ValueError:
        return None
    return candidate if candidate.is_file() else None


def _ordered_products(products: list[str], delivery: DeliveryReport) -> list[str]:
    original_positions = {product: index for index, product in enumerate(products)}
    return sorted(
        products,
        key=lambda product: (
            0 if product in delivery.tracked_rows else 1,
            delivery.tracked_rows[product].rank if product in delivery.tracked_rows else float("inf"),
            original_positions[product],
        ),
    )


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def _set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), color)


def _set_cell_margins(cell, *, top: int = 0, start: int = 90, bottom: int = 0, end: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table, color: str = BORDER_COLOR, size: str = "4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def _border(*, visible: bool) -> dict[str, str]:
    if not visible:
        return {"val": "nil"}
    return {"val": "single", "sz": "4", "space": "0", "color": BLACK}


def _set_cell_borders(cell, *, visible: bool) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        for key in list(edge.attrib):
            del edge.attrib[key]
        for key, value in _border(visible=visible).items():
            edge.set(qn(f"w:{key}"), value)


def _set_cell_width(cell, width: int) -> None:
    cell.width = Twips(width)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def _set_table_column_widths(table, widths: tuple[int, ...]) -> None:
    grid_columns = table._tbl.tblGrid.gridCol_lst
    for grid_column, width in zip(grid_columns, widths):
        grid_column.set(qn("w:w"), str(width))
    for column, width in zip(table.columns, widths):
        column.width = Twips(width)
        for cell in column.cells:
            _set_cell_width(cell, width)


def _set_spanned_table_column_widths(table, widths: tuple[int, ...]) -> None:
    for grid_column, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_column.set(qn("w:w"), str(width))
    for row in table.rows:
        column_index = 0
        for tc in row._tr.tc_lst:
            tc_pr = tc.get_or_add_tcPr()
            grid_span = tc_pr.find(qn("w:gridSpan"))
            span = int(grid_span.get(qn("w:val"))) if grid_span is not None else 1
            width = sum(widths[column_index:column_index + span])
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            column_index += span


def _set_exact_line_spacing(paragraph, points: float) -> None:
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(points)


def _set_row_height(row, points: float, *, exact: bool) -> None:
    row.height = Pt(points)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if exact else WD_ROW_HEIGHT_RULE.AT_LEAST


def _set_run_font(run, size: float, *, bold: bool = False, color: str = BLACK, strike: bool = False) -> None:
    run.font.name = LATIN_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.strike = strike
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), LATIN_FONT)
    r_fonts.set(qn("w:hAnsi"), LATIN_FONT)
    r_fonts.set(qn("w:eastAsia"), CHINESE_FONT)
    r_fonts.set(qn("w:cs"), LATIN_FONT)
    spacing = r_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        r_pr.append(spacing)
    spacing.set(qn("w:val"), LETTER_SPACING_TWENTIETHS_PT)


def _format_paragraph(
    paragraph,
    *,
    size: float = 10,
    bold: bool = False,
    alignment=None,
    line_spacing: float = 15,
) -> None:
    if alignment is not None:
        paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    _set_exact_line_spacing(paragraph, line_spacing)
    for run in paragraph.runs:
        _set_run_font(run, size, bold=bold, strike=bool(run.font.strike))


def _write_cell(cell, value: Any, *, bold: bool = False, fill: str | None = None, alignment=WD_ALIGN_PARAGRAPH.CENTER, size: float = 10) -> None:
    cell.text = _text(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(cell)
    if fill:
        _set_cell_fill(cell, fill)
    for paragraph in cell.paragraphs:
        _format_paragraph(paragraph, size=size, bold=bold, alignment=alignment)


def _style_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    for row in table.rows:
        _set_row_cant_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(cell)


def _finalize_table_rows(table, *, exact: bool = True) -> None:
    for row in table.rows:
        _set_row_height(row, 15, exact=exact)
        _set_row_cant_split(row)
        seen_cells: set[object] = set()
        for cell in row.cells:
            if cell._tc in seen_cells:
                continue
            seen_cells.add(cell._tc)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                if paragraph.alignment is None:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                _set_exact_line_spacing(paragraph, 15)


def _keep_table_together(table) -> None:
    """Keep compact tables on one page when they fit."""
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def _remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def _add_title(document: DocumentType, text: str) -> None:
    paragraph = document.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(12)
    _set_exact_line_spacing(paragraph, 20)
    _remove_paragraph_borders(paragraph)
    run = paragraph.add_run(text)
    _set_run_font(run, 16, bold=True)


def _add_heading(document: DocumentType, text: str, *, level: int = 2, keep_next: bool = True) -> None:
    paragraph = document.add_paragraph(style=f"Heading {min(level, 2)}")
    paragraph.paragraph_format.space_before = Pt(25 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = keep_next
    _set_exact_line_spacing(paragraph, 20)
    run = paragraph.add_run(text)
    _set_run_font(run, 12, bold=True)


def _add_body_paragraph(
    document: DocumentType,
    text: str,
    *,
    bold: bool = False,
    note: bool = False,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    _set_exact_line_spacing(paragraph, 15 if note else 20)
    run = paragraph.add_run(text)
    _set_run_font(run, 10 if note else 12, bold=bold, color="999999" if note else BLACK)


def _add_price_runs(paragraph, value: str, *, size: float = 12, bold: bool = False) -> None:
    text = "".join(line.strip() for line in str(value or "").replace("\r", "\n").split("\n") if line.strip())
    pattern = re.compile(r"([(（])\s*(\d+(?:\.\d+)?)\s*元\s*([)）])")
    position = 0
    for match in pattern.finditer(text):
        before = text[position:match.start()]
        if before:
            _set_run_font(paragraph.add_run(before), size, bold=bold)
        _set_run_font(paragraph.add_run(match.group(1)), size, bold=bold)
        _set_run_font(paragraph.add_run(f"{match.group(2)}元"), size, bold=bold, strike=True)
        _set_run_font(paragraph.add_run(match.group(3)), size, bold=bold)
        position = match.end()
    if position < len(text):
        _set_run_font(paragraph.add_run(text[position:]), size, bold=bold)


def _new_product_badge(product: str, new_product_dates: dict[str, Any]) -> str:
    lookup = {
        normalize_product_key(name): value
        for name, value in new_product_dates.items()
        if value
    }
    launch_date = lookup.get(normalize_product_key(product))
    return f" {launch_date.month}.{launch_date.day} NEW" if launch_date else ""


def _write_product_with_badge(
    cell,
    product: str,
    badge: str,
    *,
    bold: bool = False,
    fill: str | None = None,
) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(cell)
    if fill:
        _set_cell_fill(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    _set_exact_line_spacing(paragraph, 15)
    _set_run_font(paragraph.add_run(product), 10, bold=bold)
    if badge:
        _set_run_font(paragraph.add_run(badge), 8, bold=True, color="F28C28")


def _add_delivery_table(
    document: DocumentType,
    brand: str,
    delivery: DeliveryReport,
    tracked_products: list[str],
    new_product_dates: dict[str, Any],
) -> None:
    table = document.add_table(rows=3, cols=6)
    _style_table(table)
    widths = (4590, 1305, 1305, 1305, 1305, 630)
    _set_table_column_widths(table, widths)
    title = table.rows[0].cells[0].merge(table.rows[0].cells[-1])
    _write_cell(title, f"{brand} 美团&饿了么外卖销量表现", bold=True, fill=HEADER_FILL)
    table.rows[1].cells[0].merge(table.rows[2].cells[0])
    _write_cell(table.rows[1].cells[0], "商品名称", bold=True, fill=HEADER_FILL)
    _write_cell(table.rows[1].cells[1], "美团", bold=True, fill=HEADER_FILL)
    _write_cell(table.rows[1].cells[2], "饿了么", bold=True, fill=HEADER_FILL)
    table.rows[1].cells[3].merge(table.rows[1].cells[5])
    _write_cell(table.rows[1].cells[3], "线上合计", bold=True, fill=HEADER_FILL)
    for index, label in enumerate(("日店均销量", "日店均销量", "日店均销量", "销量占比", "排名"), 1):
        _write_cell(table.rows[2].cells[index], label, bold=True, fill=HEADER_FILL)
    for row in table.rows[:3]:
        _set_repeat_table_header(row)

    for item in delivery.display_rows:
        row = table.add_row()
        for cell, width in zip(row.cells, widths):
            _set_cell_width(cell, width)
        badge = "" if item.tracked else _new_product_badge(item.product, new_product_dates)
        fill = TRACKED_FILL if item.tracked else None
        _write_product_with_badge(row.cells[0], item.product, badge, bold=item.tracked, fill=fill)
        values = (
            _format_daily(item.meituan_daily),
            _format_daily(item.eleme_daily),
            _format_daily(item.combined_daily),
            _format_percent(item.share),
            item.rank,
        )
        for index, value in enumerate(values, 1):
            _write_cell(
                row.cells[index],
                value,
                bold=item.tracked,
                fill=fill,
                alignment=WD_ALIGN_PARAGRAPH.CENTER,
            )
    total = table.add_row()
    total_daily = delivery.meituan_total_daily + delivery.eleme_total_daily
    for cell, width in zip(total.cells, widths):
        _set_cell_width(cell, width)
    for index, value in enumerate(("总计", _format_daily(delivery.meituan_total_daily), _format_daily(delivery.eleme_total_daily), _format_daily(total_daily), "100.0%", "-")):
        _write_cell(total.cells[index], value, bold=True, fill=HEADER_FILL)
    _finalize_table_rows(table)

    all_store_products: list[str] = []
    partial_products: list[str] = []
    for product in tracked_products:
        item = delivery.tracked_rows.get(product)
        mt_stores = item.meituan_stores if item else 0
        elm_stores = item.eleme_stores if item else 0
        if delivery.meituan_stores and delivery.eleme_stores and mt_stores == delivery.meituan_stores and elm_stores == delivery.eleme_stores:
            all_store_products.append(product)
        else:
            partial_products.append(f"{product}美团在售门店 {_format_count(mt_stores)} 家，饿了么在售门店 {_format_count(elm_stores)} 家")
    clauses: list[str] = []
    if all_store_products:
        clauses.append(f"{'、'.join(all_store_products)}美团&饿了么所有门店均有售")
    clauses.extend(partial_products)
    suffix = f"。其中，{'；'.join(clauses)}。" if clauses else ""
    _add_body_paragraph(
        document,
        f"数据来源：根据{brand}美团外卖 {_format_count(delivery.meituan_stores)} 家店铺和饿了么外卖 {_format_count(delivery.eleme_stores)} 家店铺月销数据计算{suffix}",
        note=True,
    )


def _add_jd_table(
    document: DocumentType,
    brand: str,
    rows: tuple[JdSale, ...],
    total_sales: float,
    tracked_products: list[str],
    early_delisted_products: list[str],
    new_product_dates: dict[str, Any],
) -> None:
    if not rows:
        if early_delisted_products:
            _add_body_paragraph(
                document,
                f"{brand}{'、'.join(early_delisted_products)}在京东外卖提前下架，销售数据暂无法获取。",
                note=True,
            )
        else:
            _add_body_paragraph(document, f"{brand}京东外卖销售数据暂无法获取。", note=True)
        return
    tracked_keys = {normalize_product_key(product) for product in tracked_products}
    displayed = [row for row in rows if row.rank <= 20]
    displayed_keys = {normalize_product_key(row.product) for row in displayed}
    displayed.extend(row for row in rows if normalize_product_key(row.product) in tracked_keys and normalize_product_key(row.product) not in displayed_keys)

    table = document.add_table(rows=2, cols=4)
    _style_table(table)
    widths = (731, 6160, 2088, 1461)
    _set_table_column_widths(table, widths)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            _set_cell_width(cell, width)
    title = table.rows[0].cells[0].merge(table.rows[0].cells[-1])
    _write_cell(title, f"{brand} 京东外卖销量表现", bold=True, fill=HEADER_FILL)
    for index, label in enumerate(("排名", "商品名称", "京东外卖销量", "销量占比")):
        _write_cell(table.rows[1].cells[index], label, bold=True, fill=HEADER_FILL)
    for row in table.rows:
        _set_repeat_table_header(row)
    for item in displayed:
        row = table.add_row()
        key = normalize_product_key(item.product)
        tracked = key in tracked_keys
        badge = "" if tracked else _new_product_badge(item.product, new_product_dates)
        fill = TRACKED_FILL if tracked else None
        _write_cell(row.cells[0], item.rank, bold=tracked, fill=fill)
        _write_product_with_badge(row.cells[1], item.product, badge, bold=tracked, fill=fill)
        _write_cell(row.cells[2], _format_count(item.total_sales), bold=tracked, fill=fill)
        _write_cell(row.cells[3], _format_percent(item.share), bold=tracked, fill=fill)
    _finalize_table_rows(table)
    _add_body_paragraph(
        document,
        "注：京东外卖销量显示为“品牌全国门店总量”，美团/饿了么销量为“单店独立销量”。因统计口径差异，三者不可直接对比。",
        note=True,
    )
    if early_delisted_products:
        _add_body_paragraph(
            document,
            f"{brand}{'、'.join(early_delisted_products)}在京东外卖提前下架，销售数据暂无法获取。",
            note=True,
        )


def _download_image(url: str) -> BytesIO:
    if url.startswith("data:image/"):
        import base64

        return BytesIO(base64.b64decode(url.split(",", 1)[1]))
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(request, timeout=30) as response:
        data = response.read(15 * 1024 * 1024 + 1)
        if len(data) > 15 * 1024 * 1024:
            raise RuntimeError("产品图片超过 15MB")
        content_type = response.headers.get("Content-Type", "image/png").split(";", 1)[0].strip()
        if not content_type.startswith("image/"):
            raise RuntimeError(f"产品图片类型无效：{content_type}")
    return BytesIO(data)


def _add_product_info_table(document: DocumentType, product: str, info: ProductInfo | None, warnings: list[str]) -> None:
    if not info:
        _add_body_paragraph(document, "周报产品信息暂无法获取。")
        return
    table = document.add_table(rows=6, cols=2)
    _style_table(table)
    widths = (1504, 8936)
    _set_table_column_widths(table, widths)
    values = (
        ("新品名称", info.product_name or product),
        ("产品系列归属", info.series or "-"),
        ("产品卖点介绍", info.selling_point or "-"),
        ("产品价格", info.price or "-"),
        ("原料构成", info.ingredients or "-"),
        ("产品外观", ""),
    )
    for row_index, (row, (label, value)) in enumerate(zip(table.rows, values)):
        for cell, width in zip(row.cells, widths):
            _set_cell_width(cell, width)
        fill = LABEL_FILL if row_index in (0, 1) else None
        _write_cell(row.cells[0], label, bold=True, fill=fill)
        if label == "产品价格":
            row.cells[1].text = ""
            paragraph = row.cells[1].paragraphs[0]
            _add_price_runs(paragraph, value, size=10)
            _format_paragraph(paragraph, size=10, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        elif label != "产品外观":
            alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if label == "产品卖点介绍" else WD_ALIGN_PARAGRAPH.CENTER
            _write_cell(
                row.cells[1],
                value,
                bold=label == "新品名称",
                fill=fill,
                alignment=alignment,
            )
    image_cell = table.rows[-1].cells[1]
    image_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    image_paragraph = image_cell.paragraphs[0]
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if info.image_urls:
        try:
            image = _download_image(info.image_urls[0])
            image_paragraph.add_run().add_picture(image, height=Cm(4))
        except Exception as exc:
            warnings.append(f"{product}：Word 产品外观图片下载失败（{exc}）")
            _set_run_font(image_paragraph.add_run("产品外观暂无法获取"), 10, color="666666")
    else:
        _set_run_font(image_paragraph.add_run("产品外观暂无法获取"), 10, color="666666")
    _finalize_table_rows(table, exact=False)
    image_row = table.rows[-1]
    image_row.height = None
    image_row.height_rule = None
    image_paragraph.paragraph_format.line_spacing = None
    image_paragraph.paragraph_format.line_spacing_rule = None
    _keep_table_together(table)


def _add_social_summary(document: DocumentType, social: SocialReport | None) -> None:
    if not social:
        _add_body_paragraph(document, "消费者反馈统计暂无法获取。")
        return
    period = f"（{social.period}）" if social.period else ""
    if not social.total_users:
        _add_body_paragraph(document, f"上市30日{period}暂未获取到有效第三方评论。")
        return
    rate = f"，好评率为 {_format_positive_rate(social.positive_rate)}" if social.total_users >= 50 else ""
    _add_body_paragraph(
        document,
        f"上市30日{period}第三方评论共 {_format_count(social.total_users)} 条{rate}：\n"
        f"好评（{_format_count(social.positive_users)} 条）主要提及关键词：{_top_labels(social.positive_top)}；\n"
        f"差评（{_format_count(social.negative_users)} 条）主要提及关键词：{_top_labels(social.negative_top)}。",
    )


def _add_social_detail_table(document: DocumentType, social: SocialReport | None) -> None:
    if not social:
        _add_body_paragraph(document, "消费者反馈详情暂无法获取。")
        return
    table = document.add_table(rows=0, cols=6)
    _style_table(table)
    widths = (3341, 835, 3341, 835, 1462, 626)
    _set_table_column_widths(table, widths)
    kpis = iter((
        ("总计", _format_count(social.total_users)),
        ("好评用户数", _format_count(social.positive_users)),
        ("差评用户数", _format_count(social.negative_users)),
    ))

    def add_kpi(row, pair: tuple[str, str] | None) -> None:
        if pair:
            _write_cell(row.cells[4], pair[0])
            _write_cell(row.cells[5], pair[1])
        else:
            spacer = row.cells[4].merge(row.cells[5])
            _write_cell(spacer, "")
            _set_cell_borders(spacer, visible=False)

    if social.sections:
        first = table.add_row()
        first.cells[0].merge(first.cells[3])
        _write_cell(first.cells[0], social.title, bold=True, fill=HEADER_FILL)
        _write_cell(first.cells[4], "好评率")
        _write_cell(first.cells[5], _format_positive_rate(social.positive_rate) if social.total_users else "")
    for section in social.sections:
        header = table.add_row()
        labels = (f"{section.label}好评", "评论数", f"{section.label}差评", "评论数")
        for index, label in enumerate(labels):
            _write_cell(header.cells[index], label, bold=True, fill=HEADER_FILL)
        add_kpi(header, next(kpis, None))
        count = max(len(section.positive_tags), len(section.negative_tags))
        for index in range(count):
            row = table.add_row()
            positive = section.positive_tags[index] if index < len(section.positive_tags) else ("", 0)
            negative = section.negative_tags[index] if index < len(section.negative_tags) else ("", 0)
            values = (positive[0], positive[1] if positive[0] else "", negative[0], negative[1] if negative[0] else "")
            for column, value in enumerate(values):
                _write_cell(row.cells[column], value)
            add_kpi(row, next(kpis, None))
        total = table.add_row()
        for index, value in enumerate(("好评用户数", section.positive_users, "差评用户数", section.negative_users)):
            _write_cell(total.cells[index], value, bold=index in (0, 2))
        add_kpi(total, next(kpis, None))
    if not social.sections:
        for label, value in (("好评率", _format_positive_rate(social.positive_rate) if social.total_users else ""), *tuple(kpis)):
            row = table.add_row()
            row.cells[0].merge(row.cells[3])
            _write_cell(row.cells[0], social.title if len(table.rows) == 1 else "", bold=True, fill=HEADER_FILL if len(table.rows) == 1 else None)
            _write_cell(row.cells[4], label)
            _write_cell(row.cells[5], value)
    _set_spanned_table_column_widths(table, widths)
    _finalize_table_rows(table)


def build_report_docx(
    *,
    title: str,
    brand: str,
    products: list[str],
    report_date,
    product_infos: dict[str, ProductInfo],
    launch_dates: dict[str, Any],
    output_path: Path,
    configs: dict[str, dict[str, Any]],
    delivery_report: DeliveryReport,
    jd_report: tuple[tuple[JdSale, ...], float],
    social_reports: dict[str, SocialReport],
    social_paths: dict[str, Path] | None = None,
    delivery_statuses: dict[str, str] | None = None,
    jd_early_delisted_products: list[str] | None = None,
    hide_delivery_modules: bool = False,
    new_product_dates: dict[str, Any] | None = None,
) -> tuple[Path, tuple[str, ...]]:
    del report_date
    warnings: list[str] = []
    delivery_statuses = delivery_statuses or {}
    jd_early_delisted_products = jd_early_delisted_products or []
    new_product_dates = new_product_dates or {}
    social_reports = dict(social_reports)
    for product, path in (social_paths or {}).items():
        if product in social_reports:
            continue
        try:
            social_reports[product] = read_social_report(path)
        except Exception as exc:
            warnings.append(f"{product}：Word 消费者反馈统计读取失败（{exc}）")
    ordered_products = _ordered_products(products, delivery_report)

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.header_distance = Inches(0.18)
    section.footer_distance = Inches(0.2)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal_fonts = normal._element.get_or_add_rPr().get_or_add_rFonts()
    normal_fonts.set(qn("w:ascii"), LATIN_FONT)
    normal_fonts.set(qn("w:hAnsi"), LATIN_FONT)
    normal_fonts.set(qn("w:eastAsia"), CHINESE_FONT)
    normal_fonts.set(qn("w:cs"), LATIN_FONT)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = LATIN_FONT
        style.font.color.rgb = RGBColor(0, 0, 0)
        style_fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        style_fonts.set(qn("w:ascii"), LATIN_FONT)
        style_fonts.set(qn("w:hAnsi"), LATIN_FONT)
        style_fonts.set(qn("w:eastAsia"), CHINESE_FONT)
        style_fonts.set(qn("w:cs"), LATIN_FONT)

    logo_path = _safe_config_path(_text(configs.get("html_layout", {}).get("logo", {}).get("path")))
    header = section.header
    header.is_linked_to_previous = False
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_paragraph.paragraph_format.space_after = Pt(1)
    if logo_path:
        header_paragraph.add_run().add_picture(str(logo_path), width=Inches(0.8))
    else:
        warnings.append("Word 报告 Logo 不可访问，页眉已留空")

    _add_title(document, f"{brand}：{'、'.join(products)}" or title)
    launch_values = [launch_dates.get(product) for product in ordered_products if launch_dates.get(product)]
    launch_text = f"{launch_values[0].month}.{launch_values[0].day} " if launch_values and len(set(launch_values)) == 1 else ""
    _add_body_paragraph(document, f"以下是{brand} {launch_text}新品30日销量表现及消费者评论情况：", bold=True)
    intro_paragraphs = []
    for product in ordered_products:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        _set_exact_line_spacing(paragraph, 20)
        sale = delivery_report.tracked_rows.get(product)
        if sale:
            detail = f"销量排名第 {sale.rank}，线上外卖日店均 {_format_daily(sale.combined_daily)} 杯，销量占比 {_format_percent(sale.share)}"
        elif delivery_statuses.get(product):
            detail = f"{delivery_statuses[product]}，暂无月销数据"
        else:
            detail = "暂未匹配到外卖销售数据；"
        _set_run_font(paragraph.add_run(f"{product}："), 12, bold=True)
        _set_run_font(paragraph.add_run(detail), 12)
        intro_paragraphs.append(paragraph)
    if intro_paragraphs:
        intro_paragraphs[-1].paragraph_format.space_after = Pt(4)

    if not hide_delivery_modules:
        _add_delivery_table(document, brand, delivery_report, ordered_products, new_product_dates)
        _add_jd_table(document, brand, jd_report[0], jd_report[1], ordered_products, jd_early_delisted_products, new_product_dates)

    for product in ordered_products:
        info = product_infos.get(product)
        social = social_reports.get(product)
        heading = document.add_paragraph(style="Heading 1")
        heading.paragraph_format.space_before = Pt(25)
        heading.paragraph_format.space_after = Pt(5)
        heading.paragraph_format.keep_with_next = True
        _set_exact_line_spacing(heading, 20)
        _set_run_font(heading.add_run("● "), 12, bold=True)
        _set_run_font(heading.add_run(product), 12, bold=True)
        if info and info.price:
            _set_run_font(heading.add_run("（"), 12, bold=True)
            _add_price_runs(heading, info.price, size=12, bold=True)
            _set_run_font(heading.add_run("）"), 12, bold=True)
        if not info:
            warnings.append(f"{product}：Word 周报产品信息暂无法获取")
        if not social:
            warnings.append(f"{product}：Word 消费者反馈统计暂无法获取")
        _add_heading(document, "1. 产品信息")
        _add_product_info_table(document, product, info, warnings)
        _add_heading(document, "2. 新品消费者反馈汇总")
        _add_social_summary(document, social)
        _add_heading(document, "3. 消费者反馈详情")
        _add_social_detail_table(document, social)

    output_path = output_path.resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path, tuple(dict.fromkeys(warnings))
