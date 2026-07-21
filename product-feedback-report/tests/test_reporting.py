from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook

from scripts.reporting.html import (
    ProductInfo,
    SocialReport,
    SocialSection,
    _social_detail_html,
    _social_summary_html,
    build_report_html,
    combine_delivery_sales,
    delivery_report_from_metrics,
    jd_report_from_metrics,
    read_jd_sales,
    read_social_report,
    social_report_from_summaries,
)
from scripts.social.processing import PlatformSummary, build_social_feedback_workbook


class ReportingTests(unittest.TestCase):
    def _platform_workbook(self, path: Path, rows: list[tuple[str, float, float, int]]) -> Path:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(
            ["排名", "商品名称", "日店均销量", "总销量", "总销量占比", "总在售天数", "在售门店数", "在售天数", "上新日期"]
        )
        total = sum(row[2] for row in rows)
        for index, (product, daily, sales, stores) in enumerate(rows, 1):
            worksheet.append([index, product, daily, sales, sales / total if total else 0, 30 * stores, stores, 30, None])
        workbook.save(path)
        return path

    def _jd_workbook(self, path: Path) -> Path:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["排名", "商品名称", "销量加总", "在售门店数", "总销量", "总销量占比", "在售天数", "上新日期"])
        rows = [(1, "A", 100), (1, "B", 100), (3, "C", 50), (21, "D", 10)]
        total = sum(sales for _, _, sales in rows)
        for rank, product, sales in rows:
            worksheet.append([rank, product, sales, 1, sales, sales / total, 30, None])
        workbook.save(path)
        return path

    def test_delivery_combines_total_sales_and_adds_tracked_product_outside_top20(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            meituan_rows = [(f"商品{i}", 30 - i, 3000 - i * 100, 10) for i in range(1, 22)]
            eleme_rows = [(f"商品{i}", 20 - i / 2, 2000 - i * 50, 12) for i in range(1, 22)]
            meituan_rows[-1] = ("轻因·新品", 0.5, 5, 4)
            eleme_rows[-1] = ("轻因·新品", 0.4, 4, 5)
            mt = self._platform_workbook(root / "美团.xlsx", meituan_rows)
            elm = self._platform_workbook(root / "饿了么.xlsx", eleme_rows)

            report = combine_delivery_sales(mt, elm, ["新品"])

        self.assertEqual(len([row for row in report.display_rows if row.rank <= 20]), 20)
        self.assertTrue(report.display_rows[-1].tracked)
        self.assertEqual(report.tracked_rows["新品"].product, "轻因·新品")
        self.assertAlmostEqual(sum(row.share for row in report.all_rows), 1.0)
        self.assertEqual(report.meituan_stores, 10)
        self.assertEqual(report.eleme_stores, 12)

    def test_ambiguous_tracked_product_match_fails(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            mt = self._platform_workbook(
                root / "美团.xlsx", [("轻因·新品", 1, 10, 1), ("限定·新品", 1, 9, 1)]
            )
            elm = self._platform_workbook(root / "饿了么.xlsx", [("其他", 1, 8, 1)])
            with self.assertRaisesRegex(RuntimeError, "匹配到多个外卖商品"):
                combine_delivery_sales(mt, elm, ["新品"])

    def test_jd_top20_keeps_ties_and_total_uses_all_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = self._jd_workbook(Path(tmp) / "京东.xlsx")
            rows, total = read_jd_sales(path)
        self.assertEqual([row.rank for row in rows], [1, 1, 3])
        self.assertEqual(total, 260)

    def test_social_parser_aggregates_summary_and_keeps_empty_platforms_blank(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = build_social_feedback_workbook(
                brand="霸王茶姬",
                product="糯青山柠檬奶",
                start_date="6.18",
                end_date="7.18",
                summaries=[
                    PlatformSummary("weibo", "微博", (("好喝，喜欢，推荐", 9), ("清爽不腻", 5)), (("味道奇怪", 2),), 14, 2),
                    PlatformSummary("xiaohongshu", "小红书", (("清爽不腻", 3),), (("果味淡", 1),), 3, 1),
                    PlatformSummary("douyin", "抖音", (), (), 0, 0),
                    PlatformSummary("bilibili", "B站", (), (), 0, 0),
                ],
                output_dir=Path(tmp),
            )
            report = read_social_report(path)
        self.assertEqual(report.period, "6.18-7.18")
        self.assertEqual(report.total_users, 20)
        self.assertEqual(report.positive_top[0], ("清爽不腻", 8))
        self.assertNotIn("好喝，喜欢，推荐", [label for label, _ in report.positive_top])
        self.assertEqual(report.sections[-1].positive_tags, ())

    def test_html_is_self_contained_and_omits_zero_placeholders(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            mt = self._platform_workbook(root / "美团.xlsx", [("轻因·新品", 1.2, 120, 5)])
            elm = self._platform_workbook(root / "饿了么.xlsx", [("轻因·新品", 1.4, 140, 6)])
            social = build_social_feedback_workbook(
                brand="品牌",
                product="新品",
                start_date="6.18",
                end_date="7.18",
                summaries=[PlatformSummary(key, label, (), (), 0, 0) for key, label in (("weibo", "微博"), ("xiaohongshu", "小红书"), ("douyin", "抖音"), ("bilibili", "B站"))],
                output_dir=root,
            )
            output = root / "报告.html"
            build_report_html(
                title="报告",
                brand="品牌",
                products=["新品"],
                report_date=date(2026, 7, 18),
                meituan_path=mt,
                eleme_path=elm,
                jd_path=None,
                social_paths={"新品": social},
                product_infos={"新品": ProductInfo("新品", price="17元( 19 元 )")},
                launch_dates={"新品": date(2026, 6, 18)},
                output_path=output,
                configs={"html_layout": {}, "font_files": {}},
            )
            document = output.read_text(encoding="utf-8")

        self.assertIn(
            '<p class="jd-missing-note">品牌京东外卖销售数据暂无法获取。</p>',
            document,
        )
        self.assertNotIn(
            '<p class="missing-note">品牌京东外卖销售数据暂无法获取。</p>',
            document,
        )
        self.assertNotIn("暂无评论", document)
        self.assertIn(
            '<tr class="platform-total"><th>好评用户数</th><td>0</td><th>差评用户数</th><td>0</td>',
            document,
        )
        self.assertNotIn(">/</td>", document)
        self.assertNotIn("file://", document)
        self.assertNotIn("新品销量表现（30日）", document)
        self.assertIn("<strong>以下是品牌 6.18 新品30日销量表现及消费者评论情况：</strong>", document)
        self.assertIn('<tr class="table-title-row"><th colspan="6">品牌 美团&amp;饿了么外卖销量表现</th></tr>', document)
        self.assertIn(
            '<h2 class="product-title"><span aria-hidden="true">●</span>'
            '<span class="product-title-text">新品（17元('
            '<span class="price-strike">19元</span>)）</span></h2>'
            '<h3>1. 产品信息</h3>',
            document,
        )
        self.assertIn(
            '<tr class="compact"><th>产品价格</th><td>'
            '17元(<span class="price-strike">19元</span>)</td></tr>',
            document,
        )
        self.assertIn(".price-strike{text-decoration:line-through}", document)
        self.assertIn('.product-title>span[aria-hidden="true"]{font-size:16px;line-height:1}', document)
        self.assertNotIn('.product-title span{font-size:10px', document)
        self.assertIn('<tr class="name-row"><th>新品名称</th>', document)
        self.assertIn('<tr class="series-row"><th>产品系列归属</th>', document)
        self.assertIn('<tr class="ingredients-row"><th>原料构成</th>', document)
        self.assertIn("font-size:16px;line-height:26.67px", document)
        self.assertIn("font-size:13.33px;line-height:20px", document)
        self.assertIn("padding:0px 6px", document)
        self.assertIn("letter-spacing:0.02cm", document)
        self.assertIn("border:0.5px solid", document)
        self.assertIn(".table-scroll{width:100%;overflow:visible}", document)
        self.assertNotIn("overflow-x:auto", document)
        self.assertNotIn("min-width:760px", document)
        self.assertIn('class="report-viewport"', document)
        self.assertIn('class="report-scale"', document)
        self.assertIn("scaleWrap.style.transform = 'scale(' + scale + ') translateX(' + panX + 'px)'", document)
        self.assertIn("panX = dragStartPanX + dx / currentScale", document)
        self.assertIn("touch-action:pan-y pinch-zoom", document)
        self.assertIn("th,td{border-width:1.5px}", document)
        self.assertIn(".report-viewport,.report-scale{width:auto!important", document)
        self.assertIn("table{width:100%!important;max-width:100%}", document)
        self.assertIn(".total-row>*{font-weight:700;background:#D9D9D9}", document)
        self.assertIn("line-height:18px;color:#999999", document)
        self.assertIn(".source-note,.jd-note{font-size:13.33px", document)
        self.assertIn(".jd-missing-note{font-size:16px}", document)
        self.assertIn('class="product-column"><col class="platform-column"><col class="platform-column"', document)
        self.assertIn(".sales-table .product-column{width:36%}", document)
        self.assertIn(
            ".sales-table .platform-column,.sales-table .combined-column,.sales-table .share-column{width:14.5%}",
            document,
        )
        self.assertIn(".sales-table .rank-column{width:6%}", document)
        self.assertIn('<th class="rank-header">排名</th>', document)
        self.assertIn('class="rank-cell">1</td>', document)
        self.assertIn(".rank-header,.rank-cell{white-space:nowrap;word-break:keep-all", document)
        self.assertIn(".sales-table tbody td:first-child{white-space:nowrap", document)
        self.assertIn(".jd-table .rank-column{width:7%}", document)
        self.assertIn(".jd-table .product-column{width:59%}", document)
        self.assertIn("text-align:justify;text-align-last:left;margin:1px 0 20px", document)
        self.assertIn(".sales-table,.jd-table{margin-bottom:0}", document)
        self.assertIn(".platform-total th,.platform-total td{font-weight:700}", document)
        self.assertNotIn(".product-section,.product-info,.feedback-title-row{break-inside:avoid}", document)
        self.assertIn(".product-info,.feedback-title-row{break-inside:avoid}", document)
        self.assertIn(".product-info th{width:18%", document)
        self.assertIn(".product-info .ingredients-row td{text-align:center}", document)
        self.assertIn("p{margin:0}", document)
        self.assertIn("margin:20px 0 0", document)
        self.assertIn("height:4cm", document)
        self.assertNotIn("；</li>", document)
        self.assertNotIn('class="feedback-kpis"', document)

    def test_jd_table_includes_word_header_and_scope_note(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            mt = self._platform_workbook(root / "美团.xlsx", [("A", 1.2, 120, 5)])
            elm = self._platform_workbook(root / "饿了么.xlsx", [("A", 1.4, 140, 6)])
            jd = self._jd_workbook(root / "京东.xlsx")
            output = root / "报告.html"
            build_report_html(
                title="报告",
                brand="品牌",
                products=["A"],
                report_date=date(2026, 7, 18),
                meituan_path=mt,
                eleme_path=elm,
                jd_path=jd,
                social_paths={},
                product_infos={"A": ProductInfo("A")},
                launch_dates={"A": date(2026, 6, 18)},
                output_path=output,
                configs={"html_layout": {}, "font_files": {}},
            )
            document = output.read_text(encoding="utf-8")

        self.assertIn('<tr class="table-title-row"><th colspan="4">品牌 京东外卖销量表现</th></tr>', document)
        self.assertIn(
            '<tr class="tracked-row"><td class="rank-cell">1</td><td>A</td><td>100</td><td>38.5%</td></tr>',
            document,
        )
        self.assertIn('<tr class="total-row"><td colspan="2">店铺合计</td>', document)
        self.assertIn(
            '注：<strong>京东外卖销量</strong>显示为“<strong>品牌全国门店总量</strong>”，'
            '而<strong>美团/饿了么</strong>的销量为“<strong>单店独立销量</strong>”。'
            '因统计口径差异，三者不可直接对比。',
            document,
        )

    def test_social_detail_uses_word_six_column_layout_and_clean_empty_values(self) -> None:
        report = SocialReport(
            title="新品 6.18-7.17 第三方平台评价反馈",
            period="6.18-7.17",
            sections=(
                SocialSection("大众点评", (("清爽不腻", 10),), (), 10, 0),
                SocialSection("微博", (), (), 0, 0),
                SocialSection(
                    "小红书",
                    (("很长的正面评价标签用于验证内容不会被省略", 3),),
                    (("味道寡淡", 2), ("性价比低", 1)),
                    3,
                    3,
                ),
            ),
            positive_users=13,
            negative_users=3,
            positive_top=(("清爽不腻", 10),),
            negative_top=(("味道寡淡", 2),),
        )

        document = _social_detail_html(report)
        summary = _social_summary_html(report)

        self.assertEqual(document.count('<col class="'), 6)
        self.assertIn('class="feedback-title-cell" colspan="4"', document)
        self.assertIn('<th class="kpi-label">好评率</th><td class="kpi-value">81%</td>', document)
        self.assertIn('<th class="kpi-label">总计</th><td class="kpi-value">16</td>', document)
        self.assertIn('<th class="kpi-label">好评用户数</th><td class="kpi-value">13</td>', document)
        self.assertIn('<th class="kpi-label">差评用户数</th><td class="kpi-value">3</td>', document)
        self.assertNotIn("暂无评论", document)
        self.assertIn(
            '<tr class="platform-total"><th>好评用户数</th><td>0</td><th>差评用户数</th><td>0</td>',
            document,
        )
        self.assertIn('class="detail-header count-header">评论数</th>', document)
        self.assertIn("很长的正面评价标签用于验证内容不会被省略", document)
        self.assertIn('class="kpi-spacer" colspan="2"', document)
        self.assertNotIn('class="feedback-kpis"', document)
        self.assertNotIn(">/</td>", document)
        self.assertEqual(summary.count("<p"), 1)
        self.assertIn('<p class="social-summary">', summary)
        self.assertEqual(summary.count("<br>"), 2)

    def test_html_accepts_in_memory_statistics_without_intermediate_files(self) -> None:
        delivery = delivery_report_from_metrics(
            [{"product": "新品", "daily_store_avg": 1.2, "sales": 120, "stores": 5}],
            [{"product": "新品", "daily_store_avg": 1.4, "sales": 140, "stores": 6}],
            ["新品"],
        )
        jd = jd_report_from_metrics([])
        social = social_report_from_summaries(
            title="新品 6.18-7.18 第三方平台评价反馈",
            period="6.18-7.18",
            summaries=[
                PlatformSummary(key, label, (), (), 0, 0)
                for key, label in (
                    ("weibo", "微博"),
                    ("xiaohongshu", "小红书"),
                    ("douyin", "抖音"),
                    ("bilibili", "B站"),
                )
            ],
        )
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "报告.html"
            build_report_html(
                title="报告",
                brand="品牌",
                products=["新品"],
                report_date=date(2026, 7, 18),
                meituan_path=None,
                eleme_path=None,
                jd_path=None,
                social_paths={},
                product_infos={"新品": ProductInfo("新品")},
                launch_dates={"新品": date(2026, 6, 18)},
                output_path=output,
                configs={"html_layout": {}, "font_files": {}},
                delivery_report=delivery,
                jd_report=jd,
                social_report_models={"新品": social},
            )
            document = output.read_text(encoding="utf-8")
        self.assertIn("新品30日销量表现及消费者评论情况", document)
        self.assertNotIn("新品销量表现（30日）", document)
        self.assertIn("新品 6.18-7.18 第三方平台评价反馈", document)
        self.assertNotIn("品牌-新品 消费者反馈", document)


if __name__ == "__main__":
    unittest.main()
