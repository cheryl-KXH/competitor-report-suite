from __future__ import annotations

import tempfile
import unittest
import zipfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt

from scripts.reporting.html import (
    ProductInfo,
    delivery_report_from_metrics,
    jd_report_from_metrics,
    social_report_from_summaries,
)
from scripts.reporting.word import build_report_docx
from scripts.social.processing import PlatformSummary


class WordReportingTests(unittest.TestCase):
    def test_word_report_matches_html_sections_and_repeats_logo_in_header(self) -> None:
        delivery = delivery_report_from_metrics(
            [
                {"product": "蜜桃冰奶", "daily_store_avg": 3.1, "sales": 93, "stores": 20},
                {"product": "其他商品", "daily_store_avg": 2.0, "sales": 60, "stores": 20},
            ],
            [
                {"product": "蜜桃冰奶", "daily_store_avg": 2.4, "sales": 72, "stores": 18},
                {"product": "其他商品", "daily_store_avg": 1.5, "sales": 45, "stores": 18},
            ],
            ["蜜桃冰奶"],
        )
        jd = jd_report_from_metrics(
            [{"product": "蜜桃冰奶", "sales": 500}, {"product": "其他商品", "sales": 800}]
        )
        social = social_report_from_summaries(
            title="蜜桃冰奶 5.25-6.23 第三方平台评价反馈",
            period="5.25-6.23",
            summaries=[
                PlatformSummary("weibo", "微博", (("果肉多", 8),), (("太甜", 2),), 8, 2),
                PlatformSummary("xiaohongshu", "小红书", (("颜值高", 4), ("零次标签", 0)), (), 4, 0),
                PlatformSummary("bilibili", "B站", (), (), 0, 0),
            ],
        )

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "茶百道：蜜桃冰奶 20260624.docx"
            actual, warnings = build_report_docx(
                title="竞品跟踪反馈报告",
                brand="茶百道",
                products=["蜜桃冰奶"],
                report_date=date(2026, 6, 24),
                product_infos={
                    "蜜桃冰奶": ProductInfo(
                        "蜜桃冰奶",
                        series="果茶系列",
                        selling_point="桃香与奶香结合",
                        price="14元(15元)",
                        ingredients="蜜桃、牛奶",
                    )
                },
                launch_dates={"蜜桃冰奶": date(2026, 5, 25)},
                output_path=output,
                configs={
                    "html_layout": {
                        "logo": {"path": "assets/logo/report_logo.png"}
                    }
                },
                delivery_report=delivery,
                jd_report=jd,
                social_reports={"蜜桃冰奶": social},
                new_product_dates={"其他商品": date(2026, 6, 1)},
            )

            document = Document(actual)
            delivery_grid = [
                int(column.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w"))
                for column in document.tables[0]._tbl.tblGrid.gridCol_lst
            ]
            body_text = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            )
            with zipfile.ZipFile(actual) as package:
                names = set(package.namelist())
                header_xml = package.read("word/header1.xml").decode("utf-8")

            headings_1 = [
                paragraph
                for paragraph in document.paragraphs
                if paragraph.style and paragraph.style.name == "Heading 1"
            ]
            headings_2 = [
                paragraph
                for paragraph in document.paragraphs
                if paragraph.style and paragraph.style.name == "Heading 2"
            ]
            intro_items = [
                paragraph
                for paragraph in document.paragraphs
                if paragraph.style and paragraph.style.name == "List Bullet"
            ]
            source_note = next(
                paragraph
                for paragraph in document.paragraphs
                if paragraph.text.startswith("数据来源：")
            )
            jd_note = next(
                paragraph
                for paragraph in document.paragraphs
                if paragraph.text.startswith("注：京东外卖销量")
            )
            product_table = document.tables[2]
            feedback_table = document.tables[3]
            product_grid = [
                int(column.get(qn("w:w")))
                for column in product_table._tbl.tblGrid.gridCol_lst
            ]
            feedback_grid = [
                int(column.get(qn("w:w")))
                for column in feedback_table._tbl.tblGrid.gridCol_lst
            ]

        self.assertEqual(warnings, ())
        self.assertIn("茶百道：蜜桃冰奶", body_text)
        self.assertIn("茶百道 美团&饿了么外卖销量表现", body_text)
        self.assertIn("茶百道 京东外卖销量表现", body_text)
        self.assertIn("1. 产品信息", body_text)
        self.assertIn("2. 新品消费者反馈汇总", body_text)
        self.assertIn("3. 消费者反馈详情", body_text)
        self.assertIn("蜜桃冰奶 5.25-6.23 第三方平台评价反馈", body_text)
        self.assertEqual(delivery_grid, [4590, 1305, 1305, 1305, 1305, 630])
        self.assertIn("word/header1.xml", names)
        self.assertIn("word/media/image1.png", names)
        self.assertIn("w:drawing", header_xml)
        self.assertFalse(document.sections[0].different_first_page_header_footer)
        self.assertEqual([paragraph.paragraph_format.space_before for paragraph in headings_1], [Pt(25)])
        self.assertTrue(headings_2)
        self.assertTrue(all(paragraph.paragraph_format.space_before == Pt(10) for paragraph in headings_2))
        self.assertTrue(intro_items[0].runs[0].text.endswith("："))
        self.assertTrue(intro_items[0].runs[0].bold)
        self.assertEqual(intro_items[-1].paragraph_format.space_after, Pt(4))
        self.assertEqual(intro_items[0].paragraph_format.line_spacing, Pt(20))
        self.assertTrue(all(run.font.size == Pt(12) for paragraph in intro_items for run in paragraph.runs))
        self.assertTrue(all(run.font.size == Pt(12) for paragraph in headings_1 + headings_2 for run in paragraph.runs))
        self.assertEqual(source_note.paragraph_format.line_spacing_rule, WD_LINE_SPACING.EXACTLY)
        self.assertEqual(source_note.paragraph_format.line_spacing, Pt(15))
        self.assertEqual(jd_note.paragraph_format.line_spacing, Pt(15))
        self.assertTrue(all(run.font.size == Pt(10) for run in source_note.runs))
        self.assertTrue(all(str(run.font.color.rgb) == "999999" for run in source_note.runs))
        self.assertTrue(all(str(run.font.color.rgb) == "999999" for run in jd_note.runs))

        self.assertEqual(product_grid, [1504, 8936])
        self.assertEqual(feedback_grid, [3341, 835, 3341, 835, 1462, 626])
        self.assertEqual(
            [product_table.rows[0].cells[index]._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) for index in (0, 1)],
            ["E7EFFA", "E7EFFA"],
        )
        self.assertEqual(
            [product_table.rows[1].cells[index]._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) for index in (0, 1)],
            ["E7EFFA", "E7EFFA"],
        )
        self.assertEqual(
            product_table.rows[2].cells[1].paragraphs[0].alignment,
            WD_ALIGN_PARAGRAPH.JUSTIFY,
        )
        self.assertTrue(
            any(run.font.strike for run in product_table.rows[3].cells[1].paragraphs[0].runs)
        )
        self.assertIsNone(product_table.rows[-1].height)
        self.assertTrue(
            all(row.height == Pt(15) for table in (document.tables[0], feedback_table) for row in table.rows)
        )
        sample_table_paragraph = document.tables[0].rows[3].cells[1].paragraphs[0]
        self.assertEqual(sample_table_paragraph.paragraph_format.line_spacing, Pt(15))
        self.assertTrue(all(run.font.size == Pt(10) for run in sample_table_paragraph.runs))

        table_border = document.tables[0]._tbl.tblPr.find(qn("w:tblBorders")).find(qn("w:top"))
        self.assertEqual(table_border.get(qn("w:color")), "000000")
        self.assertEqual(table_border.get(qn("w:sz")), "4")
        feedback_title_fill = feedback_table.rows[0].cells[0]._tc.tcPr.find(qn("w:shd")).get(qn("w:fill"))
        self.assertEqual(feedback_title_fill, "D9D9D9")
        self.assertIsNone(feedback_table.rows[0].cells[4]._tc.tcPr.find(qn("w:shd")))
        self.assertTrue(
            all(
                row._tr.trPr is None or row._tr.trPr.find(qn("w:tblHeader")) is None
                for row in feedback_table.rows
            )
        )
        platform_total = next(
            row for row in feedback_table.rows if row.cells[0].text == "好评用户数"
        )
        self.assertIsNone(platform_total.cells[0]._tc.tcPr.find(qn("w:shd")))
        zero_comment_row = next(row for row in feedback_table.rows if row.cells[0].text == "零次标签")
        self.assertEqual(zero_comment_row.cells[1].text, "0")
        platform_totals = [
            row for row in feedback_table.rows if row.cells[0].text == "好评用户数"
        ]
        self.assertEqual(
            [cell.text for cell in platform_totals[-1].cells[:4]],
            ["好评用户数", "0", "差评用户数", "0"],
        )
        spacer = next(
            cell
            for row in feedback_table.rows
            for cell in row.cells
            if cell._tc.tcPr.find(qn("w:tcBorders")) is not None
            and cell._tc.tcPr.find(qn("w:tcBorders")).find(qn("w:right")) is not None
            and cell._tc.tcPr.find(qn("w:tcBorders")).find(qn("w:right")).get(qn("w:val")) == "nil"
        )
        self.assertEqual(
            spacer._tc.tcPr.find(qn("w:tcBorders")).find(qn("w:bottom")).get(qn("w:val")),
            "nil",
        )

        all_runs = [
            run
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
            for run in paragraph.runs
            if run.text
        ]
        self.assertTrue(all(run._element.rPr.rFonts.get(qn("w:ascii")) == "Heytea Sans Serif" for run in all_runs))
        self.assertTrue(all(run._element.rPr.rFonts.get(qn("w:eastAsia")) == "方正FW筑紫黑 简 R" for run in all_runs))
        self.assertTrue(all(run._element.rPr.find(qn("w:spacing")).get(qn("w:val")) == "11" for run in all_runs))
        badges = [run for run in all_runs if run.text.strip().endswith("NEW")]
        self.assertTrue(badges)
        self.assertTrue(all(run.font.size == Pt(8) and run.bold for run in badges))
        self.assertTrue(all(str(run.font.color.rgb) == "F28C28" for run in badges))


if __name__ == "__main__":
    unittest.main()
